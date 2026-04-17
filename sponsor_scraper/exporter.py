"""
將蒐集結果輸出為 Excel 檔案。
欄位：公司名稱、網站、Email(s)、聯絡表單網址、搜尋關鍵字、狀態
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR

# 狀態 → 背景色（ARGB）
_STATUS_COLORS = {
    "已找到Email":  "FFD4EDDA",  # 淡綠
    "已填表單":     "FFFFF3CD",  # 淡黃
    "需手動處理":   "FFFCE8CD",  # 淡橙
    "無聯絡方式":   "FFF8D7DA",  # 淡紅
    "無法連線":     "FFE2E3E5",  # 淡灰
    "填表失敗":     "FFF8D7DA",  # 淡紅
}

_HEADERS = ["公司名稱", "網站", "Email(s)", "聯絡表單網址", "搜尋關鍵字", "狀態"]
_COL_WIDTHS = [30, 40, 50, 45, 25, 14]

_THIN = Side(style="thin", color="FFB0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def export_to_excel(results: list[dict]) -> str:
    """
    輸出 Excel，回傳輸出路徑。
    results: scraper / form_filler 產出的 dict 列表
    """
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(OUTPUT_DIR, f"sponsors_{timestamp}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "贊助商聯絡清單"

    # 標頭列
    header_fill = PatternFill("solid", fgColor="FF2C5282")
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    for col_idx, (header, width) in enumerate(zip(_HEADERS, _COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 24

    # 資料列
    for row_idx, entry in enumerate(results, start=2):
        emails_str = "; ".join(entry.get("emails", []))
        status = entry.get("status", "")
        row_data = [
            entry.get("company_name", ""),
            entry.get("url", ""),
            emails_str,
            entry.get("contact_form_url") or "",
            entry.get("source_keyword", ""),
            status,
        ]

        bg_color = _STATUS_COLORS.get(status, "FFFFFFFF")
        row_fill = PatternFill("solid", fgColor=bg_color)

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = _BORDER
            # Email 欄位用等寬字型方便閱讀
            if col_idx == 3:
                cell.font = Font(name="Consolas", size=10)

        ws.row_dimensions[row_idx].height = 18

    # 凍結首列
    ws.freeze_panes = "A2"

    # 統計摘要 sheet
    _add_summary_sheet(wb, results)

    wb.save(filepath)
    print(f"\n[完成] 結果已儲存至：{filepath}")
    return filepath


def _add_summary_sheet(wb: Workbook, results: list[dict]):
    ws = wb.create_sheet("統計摘要")
    from collections import Counter
    counts = Counter(r.get("status", "未知") for r in results)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10

    title_font = Font(bold=True, size=12)
    ws.cell(row=1, column=1, value="狀態").font = title_font
    ws.cell(row=1, column=2, value="筆數").font = title_font

    for i, (status, count) in enumerate(sorted(counts.items()), start=2):
        ws.cell(row=i, column=1, value=status)
        ws.cell(row=i, column=2, value=count)

    total_row = len(counts) + 2
    ws.cell(row=total_row, column=1, value="合計").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=len(results)).font = Font(bold=True)

    email_count = sum(1 for r in results if r.get("emails"))
    ws.cell(row=total_row + 2, column=1, value="找到 Email 的網站數")
    ws.cell(row=total_row + 2, column=2, value=email_count)
    ws.cell(row=total_row + 3, column=1, value="總 Email 數")
    ws.cell(row=total_row + 3, column=2, value=sum(len(r.get("emails", [])) for r in results))
